from __future__ import annotations

import zipfile
from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy.orm import Session

from commerce.authorization import AuthorizationError, Principal
from commerce.models import (
    ChannelInventory,
    CommerceOrder,
    DataImportJob,
    DataImportRecord,
    DataImportRecordStatus,
    DataImportStatus,
    DataImportType,
    MasterProduct,
    MasterSKU,
    MembershipRole,
    Organization,
    OrganizationMembership,
    PlatformRawEvent,
    PlatformSKU,
    RawEventStatus,
    Shop,
    SKUCost,
    User,
    WarehouseInventory,
    WarehouseInventorySourceEvent,
    utcnow,
)
from commerce.services.data_import import (
    MAX_IMPORT_COLUMNS,
    MAX_IMPORT_ROWS,
    DataImportConflictError,
    DataImportNotFoundError,
    DataImportService,
    DataImportValidationError,
)
from commerce.services.ingestion import IngestionService
from commerce.services.inventory import InventoryService


@pytest.fixture
def import_context(db_session: Session) -> tuple[Principal, Principal, Principal, int, int, int]:
    first = Organization(slug="import-first", name="Import First")
    second = Organization(slug="import-second", name="Import Second")
    operator = User(email="import-operator@example.com", display_name="Operator")
    approver = User(email="import-approver@example.com", display_name="Approver")
    other = User(email="import-other@example.com", display_name="Other")
    db_session.add_all([first, second, operator, approver, other])
    db_session.flush()
    operator_membership = OrganizationMembership(
        organization_id=first.id, user_id=operator.id, role=MembershipRole.OPERATOR
    )
    approver_membership = OrganizationMembership(
        organization_id=first.id, user_id=approver.id, role=MembershipRole.APPROVER
    )
    other_membership = OrganizationMembership(
        organization_id=second.id, user_id=other.id, role=MembershipRole.OWNER
    )
    shop = Shop(
        organization_id=first.id,
        name="File Import Shop",
        platform="douyin",
        external_shop_id="file-import-shop",
        country_code="CN",
        currency="CNY",
        timezone="Asia/Shanghai",
    )
    other_shop = Shop(
        organization_id=second.id,
        name="Other Import Shop",
        platform="tiktok_shop",
        external_shop_id="other-file-import-shop",
        country_code="US",
        currency="USD",
        timezone="America/Los_Angeles",
    )
    db_session.add_all(
        [operator_membership, approver_membership, other_membership, shop, other_shop]
    )
    db_session.commit()
    return (
        Principal(operator.id, first.id, operator_membership.id, MembershipRole.OPERATOR),
        Principal(approver.id, first.id, approver_membership.id, MembershipRole.APPROVER),
        Principal(other.id, second.id, other_membership.id, MembershipRole.OWNER),
        shop.id,
        other_shop.id,
        first.id,
    )


def _catalog_csv(*, product: str = "P-100", sku: str = "SKU-100") -> bytes:
    return (
        "product_code,product_name,sku_code,sku_name,external_product_id,external_sku_id,category,title\n"
        f"{product},Product 100,{sku},SKU 100,EXT-P-100,EXT-SKU-100,Accessories,Channel title\n"
    ).encode()


def _preview_catalog(
    db_session: Session, principal: Principal, shop_id: int, *, key: str = "catalog-import-001"
) -> int:
    job = DataImportService(db_session, principal).preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key=key,
        filename="catalog.csv",
        content=_catalog_csv(),
    )
    assert job.status is DataImportStatus.PREVIEWED
    return job.id


def test_csv_imports_all_four_domains_without_platform_credentials(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    catalog_id = _preview_catalog(db_session, principal, shop_id)
    catalog = service.execute(catalog_id)
    assert catalog.status is DataImportStatus.SUCCESS
    catalog_record = service.list_records(catalog.id)[0]
    assert catalog_record.result is not None
    sku_id = int(catalog_record.result["master_sku_id"])

    order_csv = (
        b"external_order_id,platform_status,currency,total_amount,ordered_at,paid_at,"
        b"external_item_id,external_sku_id,quantity,unit_price,line_amount,title\n"
        b"ORDER-100,PAID,CNY,20.0000,2026-08-16T08:00:00+00:00,"
        b"2026-08-16T08:01:00+00:00,ITEM-100,EXT-SKU-100,2,10.0000,20.0000,Item\n"
    )
    order = service.preview(
        shop_id=shop_id,
        import_type="ORDER",
        idempotency_key="order-import-001",
        filename="orders.csv",
        content=order_csv,
    )
    assert service.execute(order.id).status is DataImportStatus.SUCCESS
    assert service.execute(order.id).status is DataImportStatus.SUCCESS

    warehouse = InventoryService(db_session, principal).create_warehouse(
        code="WH-IMPORT", name="Import Warehouse", country_code="CN", timezone="Asia/Shanghai"
    )
    inventory_csv = (
        b"inventory_kind,sku_code,warehouse_code,external_sku_id,available,reserved,incoming,damaged,source_updated_at\n"
        b"WAREHOUSE,SKU-100,WH-IMPORT,,30,2,5,1,2026-08-16T09:00:00+00:00\n"
        b"CHANNEL,SKU-100,,EXT-SKU-100,20,3,0,0,2026-08-16T09:00:00+00:00\n"
    )
    inventory = service.preview(
        shop_id=shop_id,
        import_type="INVENTORY",
        idempotency_key="inventory-import-001",
        filename="inventory.csv",
        content=inventory_csv,
    )
    assert service.execute(inventory.id).status is DataImportStatus.SUCCESS
    assert service.execute(inventory.id).status is DataImportStatus.SUCCESS

    cost_csv = (
        b"sku_code,currency,purchase_cost,packaging_cost,effective_from,source,source_reference\n"
        b"SKU-100,CNY,5.5000,0.5000,2026-08-01T00:00:00+00:00,MERCHANT_FILE,cost-sheet-august\n"
    )
    cost = service.preview(
        shop_id=shop_id,
        import_type="COST",
        idempotency_key="cost-import-001",
        filename="costs.csv",
        content=cost_csv,
    )
    assert service.execute(cost.id).status is DataImportStatus.SUCCESS
    assert service.execute(cost.id).status is DataImportStatus.SUCCESS

    records = db_session.query(DataImportRecord).all()
    assert len(records) == 5
    assert all(item.status is DataImportRecordStatus.SUCCESS for item in records)
    assert db_session.query(PlatformRawEvent).count() == 5
    assert db_session.query(CommerceOrder).count() == 1
    assert db_session.query(ChannelInventory).count() == 1
    assert db_session.query(SKUCost).count() == 1
    assert all(
        item.status is RawEventStatus.PROCESSED for item in db_session.query(PlatformRawEvent).all()
    )
    physical = db_session.query(WarehouseInventory).one()
    assert physical.warehouse_id == warehouse.id
    assert physical.master_sku_id == sku_id
    assert (physical.available, physical.reserved, physical.incoming, physical.damaged) == (
        30,
        2,
        5,
        1,
    )


def test_preview_reports_invalid_rows_and_is_idempotent(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    content = (
        b"product_code,product_name,sku_code,sku_name,external_product_id,external_sku_id\n"
        b"bad code,Product,SKU-1,SKU,EXT-P,EXT-S\n"
    )
    first = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="invalid-catalog-001",
        filename="invalid.csv",
        content=content,
    )
    assert first.status is DataImportStatus.INVALID
    assert first.invalid_records == 1
    record = service.list_records(first.id)[0]
    assert record.status is DataImportRecordStatus.INVALID
    assert record.errors
    event = db_session.get(PlatformRawEvent, record.raw_event_id)
    assert event is not None
    assert event.status is RawEventStatus.FAILED

    replay = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="different-key-same-source",
        filename="renamed.csv",
        content=content,
    )
    assert replay.id == first.id
    with pytest.raises(DataImportConflictError):
        service.preview(
            shop_id=shop_id,
            import_type="CATALOG",
            idempotency_key="invalid-catalog-001",
            filename="changed.csv",
            content=_catalog_csv(),
        )


def test_xlsx_mapping_formula_rejection_permissions_and_tenant_scope(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, approver, other, shop_id, other_shop_id, _ = import_context
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["产品编码", "产品名", "SKU编码", "SKU名", "外部产品", "外部SKU"])
    sheet.append(["P-XLSX", "XLSX Product", "SKU-XLSX", "XLSX SKU", "EXT-P-X", "EXT-S-X"])
    buffer = BytesIO()
    workbook.save(buffer)
    mapping = {
        "product_code": "产品编码",
        "product_name": "产品名",
        "sku_code": "SKU编码",
        "sku_name": "SKU名",
        "external_product_id": "外部产品",
        "external_sku_id": "外部SKU",
    }
    service = DataImportService(db_session, principal)
    job = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="xlsx-import-001",
        filename="catalog.xlsx",
        content=buffer.getvalue(),
        mapping=mapping,
    )
    assert service.execute(job.id).status is DataImportStatus.SUCCESS

    formula_book = Workbook()
    formula_sheet = formula_book.active
    formula_sheet.append(list(mapping.values()))
    formula_sheet.append(["=1+1", "Product", "SKU-F", "SKU", "EXT-P-F", "EXT-S-F"])
    formula_buffer = BytesIO()
    formula_book.save(formula_buffer)
    invalid = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="xlsx-formula-001",
        filename="formula.xlsx",
        content=formula_buffer.getvalue(),
        mapping=mapping,
    )
    assert invalid.status is DataImportStatus.INVALID
    assert invalid.errors[0]["code"] == "IMPORT_FILE_INVALID"

    with pytest.raises(AuthorizationError):
        DataImportService(db_session, approver).preview(
            shop_id=shop_id,
            import_type="CATALOG",
            idempotency_key="approver-import-001",
            filename="catalog.csv",
            content=_catalog_csv(product="P-A", sku="SKU-A"),
        )
    with pytest.raises(DataImportNotFoundError):
        DataImportService(db_session, other).get_job(job.id)
    with pytest.raises(AuthorizationError):
        service.preview(
            shop_id=other_shop_id,
            import_type="CATALOG",
            idempotency_key="cross-tenant-shop-001",
            filename="catalog.csv",
            content=_catalog_csv(product="P-CROSS", sku="SKU-CROSS"),
        )


def test_inventory_file_import_preserves_stale_event_lineage(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    service.execute(_preview_catalog(db_session, principal, shop_id, key="stale-catalog-001"))
    InventoryService(db_session, principal).create_warehouse(
        code="WH-STALE", name="Stale Warehouse", country_code="CN", timezone="Asia/Shanghai"
    )

    def inventory_content(available: int, observed_at: str) -> bytes:
        return (
            "inventory_kind,sku_code,warehouse_code,available,reserved,source_updated_at\n"
            f"WAREHOUSE,SKU-100,WH-STALE,{available},0,{observed_at}\n"
        ).encode()

    newer = service.preview(
        shop_id=shop_id,
        import_type="INVENTORY",
        idempotency_key="newer-inventory-001",
        filename="newer.csv",
        content=inventory_content(50, "2026-08-16T10:00:00+00:00"),
    )
    older = service.preview(
        shop_id=shop_id,
        import_type="INVENTORY",
        idempotency_key="older-inventory-001",
        filename="older.csv",
        content=inventory_content(5, "2026-08-16T09:00:00+00:00"),
    )
    assert service.execute(newer.id).status is DataImportStatus.SUCCESS
    assert service.execute(older.id).status is DataImportStatus.SUCCESS
    assert db_session.query(WarehouseInventory).one().available == 50
    lineage = db_session.query(WarehouseInventorySourceEvent).order_by(
        WarehouseInventorySourceEvent.raw_event_id
    )
    assert [item.applied for item in lineage] == [True, False]


def test_execute_rejects_active_lease_and_recovers_expired_job(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    job_id = _preview_catalog(db_session, principal, shop_id, key="lease-catalog-001")
    job = service.get_job(job_id)
    job.status = DataImportStatus.RUNNING
    job.lease_expires_at = utcnow() + timedelta(minutes=1)
    db_session.commit()

    with pytest.raises(DataImportConflictError, match="正在执行"):
        service.execute(job_id)

    job = service.get_job(job_id)
    job.lease_expires_at = utcnow() - timedelta(seconds=1)
    db_session.commit()
    recovered = service.execute(job_id)
    assert recovered.status is DataImportStatus.SUCCESS
    assert recovered.execution_attempts == 1
    assert recovered.lease_expires_at is None


def test_execute_rebuilds_catalog_and_cost_results_after_domain_commit(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    catalog_id = _preview_catalog(db_session, principal, shop_id, key="recovery-catalog-001")
    catalog = service.execute(catalog_id)
    catalog_record = service.list_records(catalog.id)[0]
    original_catalog_result = catalog_record.result
    catalog_record.status = DataImportRecordStatus.PROCESSING
    catalog_record.result = None
    catalog.status = DataImportStatus.RUNNING
    catalog.processed_records = 0
    catalog.finished_at = None
    catalog.lease_expires_at = utcnow() - timedelta(seconds=1)
    db_session.commit()

    recovered_catalog = service.execute(catalog.id)
    assert recovered_catalog.status is DataImportStatus.SUCCESS
    assert service.list_records(catalog.id)[0].result == original_catalog_result
    assert db_session.query(MasterProduct).count() == 1
    assert db_session.query(MasterSKU).count() == 1

    cost_csv = (
        b"sku_code,currency,purchase_cost,effective_from,source\n"
        b"SKU-100,CNY,5.5000,2026-08-01T00:00:00+00:00,MERCHANT_FILE\n"
    )
    cost = service.preview(
        shop_id=shop_id,
        import_type="COST",
        idempotency_key="recovery-cost-001",
        filename="cost.csv",
        content=cost_csv,
    )
    cost = service.execute(cost.id)
    cost_record = service.list_records(cost.id)[0]
    original_cost_result = cost_record.result
    cost_record.status = DataImportRecordStatus.PROCESSING
    cost_record.result = None
    cost.status = DataImportStatus.RUNNING
    cost.processed_records = 0
    cost.finished_at = None
    cost.lease_expires_at = utcnow() - timedelta(seconds=1)
    db_session.commit()

    recovered_cost = service.execute(cost.id)
    assert recovered_cost.status is DataImportStatus.SUCCESS
    assert service.list_records(cost.id)[0].result == original_cost_result
    assert db_session.query(SKUCost).count() == 1


def test_failed_record_can_be_retried_without_duplicate_domain_rows(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    job_id = _preview_catalog(db_session, principal, shop_id, key="retry-catalog-001")

    def fail_once(*args: object, **kwargs: object) -> dict[str, object]:
        del args, kwargs
        raise RuntimeError("simulated execution failure")

    monkeypatch.setattr(service, "_execute_record", fail_once)
    failed = service.execute(job_id)
    assert failed.status is DataImportStatus.FAILED
    assert service.list_records(job_id)[0].status is DataImportRecordStatus.FAILED

    recovered = DataImportService(db_session, principal).execute(job_id)
    assert recovered.status is DataImportStatus.SUCCESS
    assert recovered.execution_attempts == 2
    assert db_session.query(MasterProduct).count() == 1
    assert db_session.query(MasterSKU).count() == 1


def test_incomplete_preview_staging_is_never_executable(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    original = IngestionService.ingest_import_event
    calls = 0

    def fail_second_event(self: IngestionService, **kwargs: object) -> PlatformRawEvent:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("simulated preview staging failure")
        return original(self, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(IngestionService, "ingest_import_event", fail_second_event)
    content = (
        b"product_code,product_name,sku_code,sku_name,external_product_id,external_sku_id\n"
        b"P-201,Product 201,SKU-201,SKU 201,EXT-P-201,EXT-SKU-201\n"
        b"P-202,Product 202,SKU-202,SKU 202,EXT-P-202,EXT-SKU-202\n"
    )
    job = DataImportService(db_session, principal).preview(
        shop_id=shop_id,
        import_type=DataImportType.CATALOG.value,
        idempotency_key="preview-failure-001",
        filename="catalog.csv",
        content=content,
    )
    assert job.status is DataImportStatus.FAILED
    assert job.execution_attempts == 0
    assert job.errors == [{"code": "IMPORT_PREVIEW_FAILED", "message": "导入预览处理失败"}]
    with pytest.raises(DataImportValidationError, match="预览未完整生成"):
        DataImportService(db_session, principal).execute(job.id)
    assert db_session.query(MasterProduct).count() == 0


def test_order_source_time_and_channel_sku_identity_are_exact(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    service.execute(_preview_catalog(db_session, principal, shop_id, key="exact-catalog-001"))

    order_csv = (
        b"external_order_id,platform_status,currency,total_amount,ordered_at,source_updated_at,"
        b"external_item_id,external_sku_id,quantity,unit_price,line_amount\n"
        b"ORDER-EXACT,PAID,CNY,20,2026-08-16T08:00:00+00:00,"
        b"2026-08-16T08:01:00+00:00,ITEM-1,EXT-SKU-100,1,10,10\n"
        b"ORDER-EXACT,PAID,CNY,20,2026-08-16T08:00:00+00:00,"
        b"2026-08-16T08:02:00+00:00,ITEM-2,EXT-SKU-100,1,10,10\n"
    )
    order_job = service.preview(
        shop_id=shop_id,
        import_type="ORDER",
        idempotency_key="order-source-time-001",
        filename="orders.csv",
        content=order_csv,
    )
    assert order_job.status is DataImportStatus.INVALID
    assert service.list_records(order_job.id)[0].errors[0]["message"] == (
        "同一订单的订单级字段不一致"
    )

    inventory_csv = (
        b"inventory_kind,sku_code,external_sku_id,available,reserved,source_updated_at\n"
        b"CHANNEL,SKU-100,ext-sku-100,10,0,2026-08-16T09:00:00+00:00\n"
    )
    inventory_job = service.preview(
        shop_id=shop_id,
        import_type="INVENTORY",
        idempotency_key="channel-exact-sku-001",
        filename="inventory.csv",
        content=inventory_csv,
    )
    assert inventory_job.status is DataImportStatus.INVALID
    assert service.list_records(inventory_job.id)[0].errors[0]["message"] == (
        "渠道库存外部 SKU 未映射到目标主 SKU"
    )


def test_csv_parser_boundaries_and_filename_sanitization(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    bom_job = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="parser-bom-001",
        filename=r"..\..\catalog.csv",
        content=b"\xef\xbb\xbf" + _catalog_csv(product="P-BOM", sku="SKU-BOM"),
    )
    assert bom_job.status is DataImportStatus.PREVIEWED
    assert bom_job.file_name == "catalog.csv"
    assert db_session.query(MasterProduct).count() == 0
    assert db_session.query(MasterSKU).count() == 0
    assert db_session.query(PlatformSKU).count() == 0

    invalid_files = [
        ("parser-header-only-001", _catalog_csv().splitlines(keepends=True)[0]),
        ("parser-encoding-001", b"product_code\n\xff\n"),
        ("parser-malformed-001", b'product_code,product_name\n"unterminated'),
        ("parser-duplicate-header-001", b"product_code,product_code\nP-1,P-1\n"),
        ("parser-empty-header-001", b"product_code,,sku_code\nP-1,x,SKU-1\n"),
        (
            "parser-unexpected-column-001",
            _catalog_csv(product="P-EXTRA", sku="SKU-EXTRA")
            .replace(b",category,title\n", b",category,title,unexpected\n")
            .replace(b",Accessories,Channel title\n", b",Accessories,Channel title,value\n"),
        ),
        (
            "parser-sensitive-header-001",
            _catalog_csv(product="P-SECRET", sku="SKU-SECRET")
            .replace(b",category,title\n", b",category,title,password\n")
            .replace(b",Accessories,Channel title\n", b",Accessories,Channel title,do-not-log\n"),
        ),
        (
            "parser-formula-001",
            _catalog_csv(product="P-FORMULA", sku="SKU-FORMULA").replace(
                b"Product 100", b"=HYPERLINK(https://invalid.example)"
            ),
        ),
    ]
    for key, content in invalid_files:
        job = service.preview(
            shop_id=shop_id,
            import_type="CATALOG",
            idempotency_key=key,
            filename="invalid.csv",
            content=content,
        )
        assert job.status is DataImportStatus.INVALID
        assert job.processed_records == 0
        assert job.failed_records == 0

    with pytest.raises(DataImportValidationError, match="仅支持"):
        service.preview(
            shop_id=shop_id,
            import_type="CATALOG",
            idempotency_key="parser-extension-001",
            filename="catalog.txt",
            content=_catalog_csv(product="P-TXT", sku="SKU-TXT"),
        )
    with pytest.raises(DataImportValidationError, match="导入文件为空"):
        service.preview(
            shop_id=shop_id,
            import_type="CATALOG",
            idempotency_key="parser-empty-file-001",
            filename="catalog.csv",
            content=b"",
        )

    max_headers = [f"column_{index}" for index in range(MAX_IMPORT_COLUMNS)]
    headers, rows = service._read_csv(
        (",".join(max_headers) + "\n" + ",".join("x" for _ in max_headers) + "\n").encode()
    )
    assert len(headers) == MAX_IMPORT_COLUMNS
    assert len(rows) == 1
    with pytest.raises(DataImportValidationError, match="64 列"):
        service._read_csv((",".join(max_headers + ["too_many"]) + "\n").encode())

    max_rows = b"value\n" + b"x\n" * MAX_IMPORT_ROWS
    assert len(service._read_csv(max_rows)[1]) == MAX_IMPORT_ROWS
    with pytest.raises(DataImportValidationError, match="5000 行"):
        service._read_csv(max_rows + b"x\n")


def test_xlsx_rejects_malformed_multisheet_and_active_content(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)

    workbook = Workbook()
    workbook.active.append(
        [
            "product_code",
            "product_name",
            "sku_code",
            "sku_name",
            "external_product_id",
            "external_sku_id",
        ]
    )
    workbook.active.append(["P-X", "Product", "SKU-X", "SKU", "EXT-P-X", "EXT-S-X"])
    workbook.create_sheet("Unexpected")
    multisheet = BytesIO()
    workbook.save(multisheet)

    active_workbook = Workbook()
    active_workbook.active.append(["product_code"])
    active_buffer = BytesIO()
    active_workbook.save(active_buffer)
    with zipfile.ZipFile(active_buffer, mode="a") as archive:
        archive.writestr("xl/vbaProject.bin", b"not executable")

    for key, content in [
        ("xlsx-malformed-001", b"not-an-xlsx"),
        ("xlsx-multisheet-001", multisheet.getvalue()),
        ("xlsx-active-content-001", active_buffer.getvalue()),
    ]:
        job = service.preview(
            shop_id=shop_id,
            import_type="CATALOG",
            idempotency_key=key,
            filename="catalog.xlsx",
            content=content,
        )
        assert job.status is DataImportStatus.INVALID
        assert job.errors == [{"code": "IMPORT_FILE_INVALID", "message": job.errors[0]["message"]}]


def test_partial_import_retries_only_failures_and_duplicate_rows_are_idempotent(
    db_session: Session,
    import_context: tuple[Principal, Principal, Principal, int, int, int],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    principal, _, _, shop_id, _, _ = import_context
    service = DataImportService(db_session, principal)
    header = "product_code,product_name,sku_code,sku_name,external_product_id,external_sku_id\n"
    rows = []
    for index in range(100):
        rows.append(
            f"P-{index:03d},Product {index},SKU-{index:03d},SKU {index},"
            f"EXT-P-{index:03d},EXT-S-{index:03d}\n"
        )
    job = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="partial-97-of-100",
        filename="catalog.csv",
        content=(header + "".join(rows)).encode(),
    )
    assert (job.total_records, job.valid_records, job.invalid_records) == (100, 100, 0)

    original_execute_record = service._execute_record

    def fail_last_three(
        job_item: DataImportJob,
        record: DataImportRecord,
        raw_event_id: int,
        claim_token: str,
    ) -> dict[str, object]:
        payload = record.normalized_payload or {}
        if str(payload.get("sku_code", "")) in {"SKU-097", "SKU-098", "SKU-099"}:
            raise RuntimeError("simulated row execution failure")
        return original_execute_record(job_item, record, raw_event_id, claim_token)

    monkeypatch.setattr(service, "_execute_record", fail_last_three)
    completed = service.execute(job.id)
    assert completed.status is DataImportStatus.PARTIAL
    assert (completed.processed_records, completed.failed_records) == (97, 3)
    records = service.list_records(job.id)
    assert len(records) == 100
    assert sum(record.status is DataImportRecordStatus.SUCCESS for record in records) == 97
    failed = [record for record in records if record.status is DataImportRecordStatus.FAILED]
    assert len(failed) == 3
    assert all(record.error_code == "IMPORT_EXECUTION_FAILED" for record in failed)
    assert all(record.errors for record in failed)
    assert db_session.query(MasterProduct).count() == 97

    recovered = DataImportService(db_session, principal).execute(job.id)
    assert recovered.status is DataImportStatus.SUCCESS
    assert recovered.execution_attempts == 2
    assert (recovered.processed_records, recovered.failed_records) == (100, 0)
    assert db_session.query(MasterProduct).count() == 100
    assert db_session.query(MasterSKU).count() == 100

    duplicate_content = (
        header
        + "P-DUP,Duplicate,SKU-DUP,Duplicate SKU,EXT-P-DUP,EXT-S-DUP\n"
        + "P-DUP,Duplicate,SKU-DUP,Duplicate SKU,EXT-P-DUP,EXT-S-DUP\n"
    ).encode()
    duplicate = service.preview(
        shop_id=shop_id,
        import_type="CATALOG",
        idempotency_key="duplicate-source-rows-001",
        filename="duplicates.csv",
        content=duplicate_content,
    )
    duplicate = service.execute(duplicate.id)
    assert duplicate.status is DataImportStatus.SUCCESS
    assert duplicate.processed_records == 2
    assert db_session.query(MasterProduct).filter_by(code="P-DUP").count() == 1
    assert db_session.query(MasterSKU).filter_by(sku_code="SKU-DUP").count() == 1
