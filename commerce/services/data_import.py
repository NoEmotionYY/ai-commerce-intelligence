from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import secrets
import zipfile
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import and_, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from commerce.authorization import Permission, Principal, require_permission, resolve_shop
from commerce.logging import SENSITIVE_FIELD_PATTERN
from commerce.models import (
    ChannelInventorySourceEvent,
    CommerceOrderSourceEvent,
    DataImportJob,
    DataImportRecord,
    DataImportRecordStatus,
    DataImportStatus,
    DataImportType,
    MasterSKU,
    OperationLog,
    PlatformRawEvent,
    PlatformSKU,
    RawEventStatus,
    SKUCost,
    Warehouse,
    WarehouseInventorySourceEvent,
    utcnow,
)
from commerce.schemas import (
    ChannelInventorySnapshotInput,
    MasterProductCreate,
    MasterSKUCreate,
    OrderSnapshotInput,
    PlatformSKUCreate,
    SKUCostCreate,
    WarehouseInventorySnapshotInput,
)
from commerce.services.catalog import CatalogService
from commerce.services.finance import FinanceService
from commerce.services.ingestion import IngestionService
from commerce.services.inventory import InventoryService
from commerce.services.order_import import OrderImportService

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5_000
MAX_IMPORT_COLUMNS = 64
MAX_CELL_CHARACTERS = 4_096
MAX_XLSX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_XLSX_ENTRIES = 1_000
IMPORT_EXECUTION_LEASE = timedelta(minutes=15)

FIELD_SETS: dict[DataImportType, tuple[frozenset[str], frozenset[str]]] = {
    DataImportType.CATALOG: (
        frozenset(
            {
                "product_code",
                "product_name",
                "sku_code",
                "sku_name",
                "external_product_id",
                "external_sku_id",
            }
        ),
        frozenset({"category", "title"}),
    ),
    DataImportType.ORDER: (
        frozenset(
            {
                "external_order_id",
                "platform_status",
                "currency",
                "total_amount",
                "ordered_at",
                "external_item_id",
                "external_sku_id",
                "quantity",
                "unit_price",
                "line_amount",
            }
        ),
        frozenset(
            {
                "paid_at",
                "shipped_at",
                "delivered_at",
                "refunded_at",
                "settled_at",
                "source_updated_at",
                "title",
            }
        ),
    ),
    DataImportType.INVENTORY: (
        frozenset({"inventory_kind", "sku_code", "available", "reserved", "source_updated_at"}),
        frozenset({"warehouse_code", "external_sku_id", "incoming", "damaged"}),
    ),
    DataImportType.COST: (
        frozenset({"sku_code", "currency", "purchase_cost", "effective_from", "source"}),
        frozenset(
            {
                "packaging_cost",
                "domestic_shipping_cost",
                "cross_border_shipping_cost",
                "warehouse_cost",
                "other_cost",
                "effective_to",
                "source_reference",
            }
        ),
    ),
}


class DataImportConflictError(ValueError):
    pass


class DataImportNotFoundError(LookupError):
    pass


class DataImportValidationError(ValueError):
    pass


def _json_bytes(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
        "utf-8"
    )


def _digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day).isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise DataImportValidationError("表格包含非有限数值")
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    if len(text) > MAX_CELL_CHARACTERS:
        raise DataImportValidationError("单元格内容超过限制")
    if text.startswith(("=", "+", "-", "@")):
        raise DataImportValidationError("表格不允许公式样式单元格")
    return text


class DataImportService:
    def __init__(self, session: Session, principal: Principal) -> None:
        self.session = session
        self.principal = principal

    def preview(
        self,
        *,
        shop_id: int,
        import_type: str,
        idempotency_key: str,
        filename: str,
        content: bytes,
        mapping: dict[str, str] | None = None,
    ) -> DataImportJob:
        require_permission(self.principal, Permission.WRITE_COMMERCE)
        shop = resolve_shop(self.session, self.principal, shop_id)
        kind = self._import_type(import_type)
        safe_name, file_format = self._file_identity(filename)
        if not 8 <= len(idempotency_key.strip()) <= 128:
            raise DataImportValidationError("幂等键长度必须为 8 到 128")
        if not content:
            raise DataImportValidationError("导入文件为空")
        if len(content) > MAX_IMPORT_BYTES:
            raise DataImportValidationError("导入文件超过 5 MiB 限制")
        normalized_mapping = self._normalize_mapping(kind, mapping or {})
        parse_error: DataImportValidationError | None = None
        raw_records: list[tuple[int, dict[str, object]]] = []
        try:
            headers, physical_rows = self._read_file(content, file_format)
            effective_mapping = self._resolve_mapping(kind, normalized_mapping, headers)
            raw_records = self._logical_records(kind, physical_rows, effective_mapping)
            if not raw_records:
                raise DataImportValidationError("导入文件没有数据行")
        except DataImportValidationError as exc:
            parse_error = exc
            effective_mapping = normalized_mapping
        content_hash = _digest(content)
        mapping_hash = _digest(_json_bytes(effective_mapping))
        source_identity_hash = _digest(
            _json_bytes(
                {
                    "shop_id": shop.id,
                    "import_type": kind.value,
                    "content_hash": content_hash,
                    "mapping_hash": mapping_hash,
                }
            )
        )
        request_hash = _digest(
            _json_bytes(
                {
                    "shop_id": shop.id,
                    "import_type": kind.value,
                    "content_hash": content_hash,
                    "mapping": effective_mapping,
                }
            )
        )
        key_hash = _digest(idempotency_key.strip().encode("utf-8"))
        existing = self.session.scalar(
            select(DataImportJob).where(
                DataImportJob.organization_id == self.principal.organization_id,
                DataImportJob.idempotency_key_hash == key_hash,
            )
        )
        if existing is not None:
            if existing.request_hash == request_hash:
                return existing
            raise DataImportConflictError("幂等键已用于不同的导入请求")
        existing_source = self.session.scalar(
            select(DataImportJob).where(
                DataImportJob.organization_id == self.principal.organization_id,
                DataImportJob.shop_id == shop.id,
                DataImportJob.source_identity_hash == source_identity_hash,
            )
        )
        if existing_source is not None:
            return existing_source

        job = DataImportJob(
            organization_id=self.principal.organization_id,
            shop_id=shop.id,
            import_type=kind,
            file_name=safe_name,
            file_format=file_format,
            content_hash=content_hash,
            mapping=effective_mapping,
            mapping_hash=mapping_hash,
            source_identity_hash=source_identity_hash,
            idempotency_key_hash=key_hash,
            request_hash=request_hash,
            status=DataImportStatus.PREVIEWED,
            total_records=0,
            valid_records=0,
            invalid_records=0,
            processed_records=0,
            failed_records=0,
            execution_attempts=0,
            errors=[],
            created_by_user_id=self.principal.user_id,
        )
        self.session.add(job)
        try:
            self.session.commit()
        except IntegrityError as exc:
            self.session.rollback()
            concurrent = self.session.scalar(
                select(DataImportJob).where(
                    DataImportJob.organization_id == self.principal.organization_id,
                    DataImportJob.shop_id == shop.id,
                    DataImportJob.source_identity_hash == source_identity_hash,
                )
            )
            if concurrent is not None and concurrent.request_hash == request_hash:
                return concurrent
            raise DataImportConflictError("导入请求发生并发冲突") from exc

        if parse_error is not None:
            job = self._job(job.id, for_update=True)
            job.status = DataImportStatus.INVALID
            job.errors = [{"code": "IMPORT_FILE_INVALID", "message": str(parse_error)}]
            self._audit("data_import.preview", {"import_job_id": job.id, "status": "INVALID"})
            self.session.commit()
            return job
        try:
            self._stage_records(job, raw_records, effective_mapping)
        except Exception:
            self.session.rollback()
            job = self._job(job.id, for_update=True)
            records = self._records(job.id)
            job.total_records = len(records)
            job.valid_records = sum(item.status is DataImportRecordStatus.VALID for item in records)
            job.invalid_records = sum(
                item.status is DataImportRecordStatus.INVALID for item in records
            )
            job.status = DataImportStatus.FAILED
            job.errors = [{"code": "IMPORT_PREVIEW_FAILED", "message": "导入预览处理失败"}]
            self._audit(
                "data_import.preview",
                {"import_job_id": job.id, "shop_id": job.shop_id, "status": "FAILED"},
            )
            self.session.commit()
            return job

        job = self._job(job.id, for_update=True)
        records = self._records(job.id)
        job.total_records = len(records)
        job.valid_records = sum(item.status is DataImportRecordStatus.VALID for item in records)
        job.invalid_records = sum(item.status is DataImportRecordStatus.INVALID for item in records)
        job.status = DataImportStatus.PREVIEWED if job.valid_records else DataImportStatus.INVALID
        self._audit(
            "data_import.preview",
            {
                "import_job_id": job.id,
                "shop_id": job.shop_id,
                "import_type": job.import_type.value,
                "content_hash": job.content_hash,
                "total_records": job.total_records,
                "valid_records": job.valid_records,
                "invalid_records": job.invalid_records,
            },
        )
        self.session.commit()
        return job

    def execute(self, import_job_id: int) -> DataImportJob:
        require_permission(self.principal, Permission.WRITE_COMMERCE)
        job = self._job(import_job_id, for_update=True)
        if job.status is DataImportStatus.SUCCESS:
            return job
        now = utcnow()
        if (
            job.status is DataImportStatus.RUNNING
            and job.lease_expires_at is not None
            and job.lease_expires_at > now
        ):
            raise DataImportConflictError("导入任务正在执行")
        if job.status is DataImportStatus.FAILED and job.execution_attempts == 0:
            raise DataImportValidationError("导入预览未完整生成，不能执行")
        if job.status is DataImportStatus.INVALID or job.valid_records == 0:
            raise DataImportValidationError("导入预览没有可执行记录")
        recovered = job.status is DataImportStatus.RUNNING
        job.status = DataImportStatus.RUNNING
        job.execution_attempts += 1
        job.started_at = now
        job.finished_at = None
        job.lease_expires_at = now + IMPORT_EXECUTION_LEASE
        self._audit(
            "data_import.execute_start",
            {
                "import_job_id": job.id,
                "attempt": job.execution_attempts,
                "recovered": recovered,
            },
        )
        self.session.commit()

        for record in self._records(job.id):
            if record.status in {
                DataImportRecordStatus.INVALID,
                DataImportRecordStatus.SUCCESS,
            }:
                continue
            claim_token = secrets.token_urlsafe(32)
            try:
                recovered_result = self._completed_record_result(job, record)
                if recovered_result is not None:
                    self._mark_record_success(record.id, recovered_result)
                    self._extend_execution_lease(job.id)
                    continue
                if record.status in {
                    DataImportRecordStatus.FAILED,
                    DataImportRecordStatus.PROCESSING,
                }:
                    IngestionService(self.session, self.principal).retry_import_event(record.id)
                    record = self._record(record.id)
                if record.status is not DataImportRecordStatus.VALID:
                    raise DataImportConflictError("导入记录状态无法恢复")
                event = IngestionService(self.session, self.principal).begin_import_event(
                    record.id, claim_token=claim_token
                )
                result = self._execute_record(job, record, event.id, claim_token)
                self._mark_record_success(record.id, result)
            except Exception as exc:
                self.session.rollback()
                self._fail_record(record.id, claim_token, exc)
            self._extend_execution_lease(job.id)

        job = self._job(job.id, for_update=True)
        records = self._records(job.id)
        job.processed_records = sum(
            item.status is DataImportRecordStatus.SUCCESS for item in records
        )
        job.failed_records = sum(item.status is DataImportRecordStatus.FAILED for item in records)
        has_errors = bool(job.invalid_records or job.failed_records)
        if job.processed_records and has_errors:
            job.status = DataImportStatus.PARTIAL
        elif job.processed_records == job.valid_records and not has_errors:
            job.status = DataImportStatus.SUCCESS
        else:
            job.status = DataImportStatus.FAILED
        job.finished_at = utcnow()
        job.lease_expires_at = None
        self._audit(
            "data_import.execute_finish",
            {
                "import_job_id": job.id,
                "status": job.status.value,
                "processed_records": job.processed_records,
                "failed_records": job.failed_records,
            },
        )
        self.session.commit()
        return job

    def _mark_record_success(self, record_id: int, result: dict[str, object]) -> None:
        record = self._record(record_id, for_update=True)
        record.status = DataImportRecordStatus.SUCCESS
        record.result = result
        record.error_code = None
        record.errors = []
        self.session.commit()

    def _extend_execution_lease(self, import_job_id: int) -> None:
        job = self._job(import_job_id, for_update=True)
        if job.status is DataImportStatus.RUNNING:
            job.lease_expires_at = utcnow() + IMPORT_EXECUTION_LEASE
            self.session.commit()

    def _completed_record_result(
        self, job: DataImportJob, record: DataImportRecord
    ) -> dict[str, object] | None:
        if record.raw_event_id is None:
            return None
        event = self.session.scalar(
            select(PlatformRawEvent).where(
                PlatformRawEvent.id == record.raw_event_id,
                PlatformRawEvent.organization_id == self.principal.organization_id,
                PlatformRawEvent.shop_id == job.shop_id,
            )
        )
        if event is None:
            raise DataImportConflictError("导入记录的原始事件不存在")
        if event.status is not RawEventStatus.PROCESSED:
            return None
        if job.import_type is DataImportType.ORDER:
            source = self.session.scalar(
                select(CommerceOrderSourceEvent).where(
                    CommerceOrderSourceEvent.organization_id == self.principal.organization_id,
                    CommerceOrderSourceEvent.shop_id == job.shop_id,
                    CommerceOrderSourceEvent.raw_event_id == event.id,
                )
            )
            if source is None:
                raise DataImportConflictError("已完成的订单导入缺少领域证据")
            return {"commerce_order_id": source.order_id}
        if job.import_type is DataImportType.INVENTORY:
            warehouse_source = self.session.scalar(
                select(WarehouseInventorySourceEvent).where(
                    WarehouseInventorySourceEvent.organization_id == self.principal.organization_id,
                    WarehouseInventorySourceEvent.raw_event_id == event.id,
                )
            )
            if warehouse_source is not None:
                return {"warehouse_inventory_id": warehouse_source.warehouse_inventory_id}
            channel_source = self.session.scalar(
                select(ChannelInventorySourceEvent).where(
                    ChannelInventorySourceEvent.organization_id == self.principal.organization_id,
                    ChannelInventorySourceEvent.shop_id == job.shop_id,
                    ChannelInventorySourceEvent.raw_event_id == event.id,
                )
            )
            if channel_source is None:
                raise DataImportConflictError("已完成的库存导入缺少领域证据")
            return {"channel_inventory_id": channel_source.channel_inventory_id}
        payload = record.normalized_payload
        if payload is None:
            raise DataImportConflictError("已完成的导入记录缺少规范化证据")
        if job.import_type is DataImportType.CATALOG:
            sku = self.session.scalar(
                select(MasterSKU).where(
                    MasterSKU.organization_id == self.principal.organization_id,
                    MasterSKU.sku_code == str(payload["sku_code"]),
                )
            )
            external_sku_id = str(payload["external_sku_id"])
            mapping = self.session.scalar(
                select(PlatformSKU).where(
                    PlatformSKU.organization_id == self.principal.organization_id,
                    PlatformSKU.shop_id == job.shop_id,
                    PlatformSKU.master_sku_id == (sku.id if sku is not None else -1),
                    PlatformSKU.external_sku_key
                    == hashlib.sha256(external_sku_id.encode("utf-8")).hexdigest(),
                    PlatformSKU.external_sku_id == external_sku_id,
                )
            )
            if sku is None or mapping is None:
                raise DataImportConflictError("已完成的商品导入缺少领域证据")
            return {
                "master_product_id": sku.master_product_id,
                "master_sku_id": sku.id,
                "platform_sku_id": mapping.id,
            }
        model = SKUCostCreate.model_validate(payload)
        cost = self.session.scalar(
            select(SKUCost).where(
                SKUCost.organization_id == self.principal.organization_id,
                SKUCost.master_sku_id == model.master_sku_id,
                SKUCost.effective_from == model.effective_from,
            )
        )
        if cost is None or not FinanceService.cost_matches(cost, model):
            raise DataImportConflictError("已完成的成本导入缺少领域证据")
        return {"sku_cost_id": cost.id}

    def get_job(self, import_job_id: int) -> DataImportJob:
        require_permission(self.principal, Permission.READ_COMMERCE)
        return self._job(import_job_id)

    def list_jobs(
        self,
        *,
        shop_id: int | None = None,
        after_id: int = 0,
        limit: int = 50,
    ) -> list[DataImportJob]:
        require_permission(self.principal, Permission.READ_COMMERCE)
        if after_id < 0 or not 1 <= limit <= 200:
            raise DataImportValidationError("导入任务分页参数无效")
        statement = select(DataImportJob).where(
            DataImportJob.organization_id == self.principal.organization_id,
            DataImportJob.id > after_id,
        )
        if shop_id is not None:
            resolve_shop(self.session, self.principal, shop_id, require_active=False)
            statement = statement.where(DataImportJob.shop_id == shop_id)
        return list(self.session.scalars(statement.order_by(DataImportJob.id).limit(limit)))

    def list_records(self, import_job_id: int) -> list[DataImportRecord]:
        require_permission(self.principal, Permission.READ_COMMERCE)
        self._job(import_job_id)
        return self._records(import_job_id)

    def _stage_records(
        self,
        job: DataImportJob,
        raw_records: list[tuple[int, dict[str, object]]],
        mapping: dict[str, str],
    ) -> None:
        ingestion = IngestionService(self.session, self.principal)
        for row_number, raw_values in raw_records:
            record_key = _digest(_json_bytes({"row_number": row_number, "raw_values": raw_values}))
            normalized: dict[str, object] | None = None
            errors: list[dict[str, object]] = []
            status = DataImportRecordStatus.VALID
            try:
                normalized = self._validate_record(job, raw_values, mapping)
            except (DataImportValidationError, ValidationError) as exc:
                status = DataImportRecordStatus.INVALID
                errors = self._validation_errors(exc)
            record = DataImportRecord(
                organization_id=self.principal.organization_id,
                shop_id=job.shop_id,
                import_job_id=job.id,
                row_number=row_number,
                record_key=record_key,
                raw_values=raw_values,
                normalized_payload=normalized,
                errors=errors,
                status=status,
            )
            self.session.add(record)
            self.session.commit()
            event_type = self._event_type(job.import_type, raw_values, mapping)
            event = ingestion.ingest_import_event(
                import_record_id=record.id,
                event_type=event_type,
                external_event_id=f"FILE:{job.source_identity_hash[:24]}:{record_key[:24]}",
                payload={
                    "source": "FILE_IMPORT",
                    "import_job_id": job.id,
                    "record": raw_values,
                },
                occurred_at=self._source_time(job.import_type, raw_values, mapping),
            )
            del event
            if status is DataImportRecordStatus.INVALID:
                ingestion.reject_import_event(record.id, error_code="IMPORT_VALIDATION_FAILED")

    def _execute_record(
        self, job: DataImportJob, record: DataImportRecord, raw_event_id: int, claim_token: str
    ) -> dict[str, object]:
        payload = record.normalized_payload
        if payload is None:
            raise DataImportValidationError("导入记录缺少规范化结果")
        if job.import_type is DataImportType.CATALOG:
            catalog_service = CatalogService(self.session, self.principal)
            product = catalog_service.create_product(
                code=str(payload["product_code"]),
                name=str(payload["product_name"]),
                category=self._optional_text(payload.get("category")),
            )
            sku = catalog_service.create_sku(
                master_product_id=product.id,
                sku_code=str(payload["sku_code"]),
                name=str(payload["sku_name"]),
            )
            platform_sku = catalog_service.map_platform_sku(
                shop_id=job.shop_id,
                master_sku_id=sku.id,
                external_product_id=str(payload["external_product_id"]),
                external_sku_id=str(payload["external_sku_id"]),
                title=self._optional_text(payload.get("title")),
            )
            IngestionService(self.session, self.principal).complete_import_event(
                record.id, claim_token=claim_token
            )
            return {
                "master_product_id": product.id,
                "master_sku_id": sku.id,
                "platform_sku_id": platform_sku.id,
            }
        if job.import_type is DataImportType.ORDER:
            order = OrderImportService(self.session, self.principal).import_snapshot(
                raw_event_id=raw_event_id,
                claim_token=claim_token,
                import_record_id=record.id,
                snapshot=OrderSnapshotInput.model_validate(payload),
            )
            return {"commerce_order_id": order.id}
        if job.import_type is DataImportType.INVENTORY:
            inventory_service = InventoryService(self.session, self.principal)
            if payload["inventory_kind"] == "WAREHOUSE":
                warehouse_inventory = inventory_service.reconcile_warehouse_snapshot(
                    raw_event_id=raw_event_id,
                    claim_token=claim_token,
                    import_record_id=record.id,
                    snapshot=WarehouseInventorySnapshotInput.model_validate(payload["snapshot"]),
                )
                return {"warehouse_inventory_id": warehouse_inventory.id}
            channel_inventory = inventory_service.reconcile_channel_snapshot(
                raw_event_id=raw_event_id,
                claim_token=claim_token,
                import_record_id=record.id,
                snapshot=ChannelInventorySnapshotInput.model_validate(payload["snapshot"]),
            )
            return {"channel_inventory_id": channel_inventory.id}
        cost = FinanceService(self.session, self.principal).create_sku_cost(
            SKUCostCreate.model_validate(payload)
        )
        IngestionService(self.session, self.principal).complete_import_event(
            record.id, claim_token=claim_token
        )
        return {"sku_cost_id": cost.id}

    def _fail_record(self, record_id: int, claim_token: str, exc: Exception) -> None:
        safe_types = (
            DataImportConflictError,
            DataImportValidationError,
            ValidationError,
            ValueError,
            LookupError,
            PermissionError,
        )
        message = str(exc) if isinstance(exc, safe_types) else "导入执行失败"
        try:
            IngestionService(self.session, self.principal).fail_import_event(
                record_id, claim_token=claim_token, error_code="IMPORT_EXECUTION_FAILED"
            )
        except Exception:
            self.session.rollback()
        record = self._record(record_id, for_update=True)
        record.status = DataImportRecordStatus.FAILED
        record.error_code = "IMPORT_EXECUTION_FAILED"
        record.errors = [{"code": "IMPORT_EXECUTION_FAILED", "message": message[:500]}]
        self.session.commit()

    def _validate_record(
        self, job: DataImportJob, raw_values: dict[str, object], mapping: dict[str, str]
    ) -> dict[str, object]:
        if job.import_type is DataImportType.ORDER:
            raw_rows = raw_values.get("rows")
            if not isinstance(raw_rows, list):
                raise DataImportValidationError("订单导入记录结构无效")
            logical_rows = [self._apply_mapping(row, mapping) for row in raw_rows]
            return self._validate_order(logical_rows, job.shop_id)
        logical = self._apply_mapping(raw_values, mapping)
        if job.import_type is DataImportType.CATALOG:
            return self._validate_catalog(logical, job.shop_id)
        if job.import_type is DataImportType.INVENTORY:
            return self._validate_inventory(logical, job.shop_id)
        return self._validate_cost(logical)

    def _validate_catalog(self, values: dict[str, str], shop_id: int) -> dict[str, object]:
        product = MasterProductCreate(
            code=values["product_code"],
            name=values["product_name"],
            category=self._optional_text(values.get("category")),
        )
        sku = MasterSKUCreate(sku_code=values["sku_code"], name=values["sku_name"])
        platform = PlatformSKUCreate(
            shop_id=shop_id,
            master_sku_id=1,
            external_product_id=values["external_product_id"],
            external_sku_id=values["external_sku_id"],
            title=self._optional_text(values.get("title")),
        )
        return {
            "product_code": product.code,
            "product_name": product.name,
            "category": product.category,
            "sku_code": sku.sku_code,
            "sku_name": sku.name,
            "external_product_id": platform.external_product_id,
            "external_sku_id": platform.external_sku_id,
            "title": platform.title,
        }

    def _validate_order(self, rows: list[dict[str, str]], shop_id: int) -> dict[str, object]:
        if not rows or len(rows) > 1_000:
            raise DataImportValidationError("单个订单明细数量必须为 1 到 1000")
        snapshot_fields = [
            "external_order_id",
            "platform_status",
            "currency",
            "total_amount",
            "ordered_at",
            "paid_at",
            "shipped_at",
            "delivered_at",
            "refunded_at",
            "settled_at",
        ]
        order_fields = [*snapshot_fields, "source_updated_at"]
        first = rows[0]
        for row in rows[1:]:
            if any(row.get(field, "") != first.get(field, "") for field in order_fields):
                raise DataImportValidationError("同一订单的订单级字段不一致")
        payload: dict[str, object] = {
            field: self._optional_text(first.get(field)) for field in snapshot_fields
        }
        payload.update(
            {
                "external_order_id": first["external_order_id"],
                "platform_status": first["platform_status"],
                "currency": first["currency"],
                "total_amount": first["total_amount"],
                "ordered_at": first["ordered_at"],
                "items": [
                    {
                        "external_item_id": row["external_item_id"],
                        "external_sku_id": row["external_sku_id"],
                        "quantity": row["quantity"],
                        "unit_price": row["unit_price"],
                        "line_amount": row["line_amount"],
                        "title": self._optional_text(row.get("title")),
                    }
                    for row in rows
                ],
            }
        )
        model = OrderSnapshotInput.model_validate(payload)
        external_skus = {item.external_sku_id for item in model.items}
        mappings = list(
            self.session.scalars(
                select(PlatformSKU).where(
                    PlatformSKU.organization_id == self.principal.organization_id,
                    PlatformSKU.shop_id == shop_id,
                    PlatformSKU.external_sku_id.in_(external_skus),
                    PlatformSKU.active.is_(True),
                )
            )
        )
        if {item.external_sku_id for item in mappings} != external_skus:
            raise DataImportValidationError("订单包含未映射的外部 SKU")
        return model.model_dump(mode="json")

    def _validate_inventory(self, values: dict[str, str], shop_id: int) -> dict[str, object]:
        kind = values["inventory_kind"].strip().upper()
        sku = self._sku_by_code(values["sku_code"])
        if kind == "WAREHOUSE":
            warehouse_code = values.get("warehouse_code", "").strip().upper()
            warehouse = self.session.scalar(
                select(Warehouse).where(
                    Warehouse.organization_id == self.principal.organization_id,
                    Warehouse.code == warehouse_code,
                    Warehouse.active.is_(True),
                )
            )
            if warehouse is None:
                raise DataImportValidationError("仓库编码不存在或不可用")
            warehouse_snapshot = WarehouseInventorySnapshotInput(
                warehouse_id=warehouse.id,
                master_sku_id=sku.id,
                available=values["available"],
                reserved=values["reserved"],
                incoming=values.get("incoming") or "0",
                damaged=values.get("damaged") or "0",
            )
        elif kind == "CHANNEL":
            external_sku_id = values.get("external_sku_id", "")
            external_sku_key = hashlib.sha256(external_sku_id.encode("utf-8")).hexdigest()
            mapping = self.session.scalar(
                select(PlatformSKU).where(
                    PlatformSKU.organization_id == self.principal.organization_id,
                    PlatformSKU.shop_id == shop_id,
                    PlatformSKU.external_sku_key == external_sku_key,
                    PlatformSKU.external_sku_id == external_sku_id,
                    PlatformSKU.master_sku_id == sku.id,
                    PlatformSKU.active.is_(True),
                )
            )
            if mapping is None:
                raise DataImportValidationError("渠道库存外部 SKU 未映射到目标主 SKU")
            channel_snapshot = ChannelInventorySnapshotInput(
                platform_sku_id=mapping.id,
                available=values["available"],
                reserved=values["reserved"],
            )
        else:
            raise DataImportValidationError("inventory_kind 必须为 WAREHOUSE 或 CHANNEL")
        source_time = self._aware_datetime(values["source_updated_at"], "source_updated_at")
        if kind == "WAREHOUSE":
            return {
                "inventory_kind": kind,
                "source_updated_at": source_time.isoformat(),
                "snapshot": warehouse_snapshot.model_dump(mode="json"),
            }
        return {
            "inventory_kind": kind,
            "source_updated_at": source_time.isoformat(),
            "snapshot": channel_snapshot.model_dump(mode="json"),
        }

    def _validate_cost(self, values: dict[str, str]) -> dict[str, object]:
        sku = self._sku_by_code(values["sku_code"])
        model = SKUCostCreate(
            master_sku_id=sku.id,
            currency=values["currency"],
            purchase_cost=values["purchase_cost"],
            packaging_cost=values.get("packaging_cost") or "0",
            domestic_shipping_cost=values.get("domestic_shipping_cost") or "0",
            cross_border_shipping_cost=values.get("cross_border_shipping_cost") or "0",
            warehouse_cost=values.get("warehouse_cost") or "0",
            other_cost=values.get("other_cost") or "0",
            effective_from=values["effective_from"],
            effective_to=self._optional_text(values.get("effective_to")),
            source=values["source"],
            source_reference=self._optional_text(values.get("source_reference")),
        )
        overlap = [
            SKUCost.organization_id == self.principal.organization_id,
            SKUCost.master_sku_id == sku.id,
            or_(SKUCost.effective_to.is_(None), SKUCost.effective_to > model.effective_from),
        ]
        if model.effective_to is not None:
            overlap.append(SKUCost.effective_from < model.effective_to)
        existing = self.session.scalar(select(SKUCost).where(and_(*overlap)).limit(1))
        if existing is not None and not FinanceService.cost_matches(existing, model):
            raise DataImportValidationError("SKU 成本生效区间与现有记录重叠")
        return model.model_dump(mode="json")

    def _read_file(
        self, content: bytes, file_format: str
    ) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
        if file_format == "csv":
            return self._read_csv(content)
        return self._read_xlsx(content)

    def _read_csv(self, content: bytes) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DataImportValidationError("CSV 必须使用 UTF-8 编码") from exc
        if "\x00" in text:
            raise DataImportValidationError("CSV 包含 NUL 字符")
        try:
            rows: list[list[object]] = []
            for row_index, row in enumerate(csv.reader(io.StringIO(text, newline=""), strict=True)):
                if row_index > MAX_IMPORT_ROWS:
                    raise DataImportValidationError("导入数据超过 5000 行限制")
                if len(row) > MAX_IMPORT_COLUMNS:
                    raise DataImportValidationError("导入列数超过 64 列限制")
                rows.append(list(row))
        except csv.Error as exc:
            raise DataImportValidationError("CSV 结构无效") from exc
        return self._tabular_rows(rows)

    def _read_xlsx(self, content: bytes) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > MAX_XLSX_ENTRIES:
                    raise DataImportValidationError("XLSX 压缩条目过多")
                if sum(item.file_size for item in infos) > MAX_XLSX_EXPANDED_BYTES:
                    raise DataImportValidationError("XLSX 解压后超过限制")
                names = [item.filename.replace("\\", "/").lower() for item in infos]
                if len(names) != len(set(names)):
                    raise DataImportValidationError("XLSX 包含重复压缩条目")
                if any(
                    name.endswith("/vbaproject.bin")
                    or name.startswith("xl/externallinks/")
                    or name == "xl/connections.xml"
                    for name in names
                ):
                    raise DataImportValidationError("XLSX 不允许宏、外部链接或数据连接")
        except zipfile.BadZipFile as exc:
            raise DataImportValidationError("XLSX 文件结构无效") from exc
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as exc:
            raise DataImportValidationError("XLSX 无法读取") from exc
        try:
            if len(workbook.worksheets) != 1:
                raise DataImportValidationError("XLSX 必须仅包含一个工作表")
            if (workbook.active.max_row or 0) > MAX_IMPORT_ROWS + 1:
                raise DataImportValidationError("导入数据超过 5000 行限制")
            if (workbook.active.max_column or 0) > MAX_IMPORT_COLUMNS:
                raise DataImportValidationError("导入列数超过 64 列限制")
            rows: list[list[object]] = []
            try:
                for sheet_row in workbook.active.iter_rows():
                    values: list[object] = []
                    for cell in sheet_row:
                        if cell.data_type == "f":
                            raise DataImportValidationError("XLSX 不允许公式单元格")
                        values.append(cell.value)
                    rows.append(values)
                    if len(rows) > MAX_IMPORT_ROWS + 1:
                        raise DataImportValidationError("导入数据超过 5000 行限制")
            except DataImportValidationError:
                raise
            except Exception as exc:
                raise DataImportValidationError("XLSX 工作表结构无效") from exc
            return self._tabular_rows(rows)
        finally:
            workbook.close()

    def _tabular_rows(
        self, rows: list[list[object]]
    ) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
        while rows and not any(_cell_text(value) for value in rows[-1]):
            rows.pop()
        if not rows:
            raise DataImportValidationError("表格为空")
        headers = [_cell_text(value) for value in rows[0]]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not item for item in headers):
            raise DataImportValidationError("表头不得为空")
        if len(headers) > MAX_IMPORT_COLUMNS:
            raise DataImportValidationError("导入列数超过 64 列限制")
        if len(set(headers)) != len(headers):
            raise DataImportValidationError("表头不得重复")
        if len(rows) - 1 > MAX_IMPORT_ROWS:
            raise DataImportValidationError("导入数据超过 5000 行限制")
        result: list[tuple[int, dict[str, object]]] = []
        for row_number, row in enumerate(rows[1:], start=2):
            values = [_cell_text(value) for value in row[: len(headers)]]
            values.extend([""] * (len(headers) - len(values)))
            if not any(values):
                continue
            if len(row) > len(headers) and any(_cell_text(value) for value in row[len(headers) :]):
                raise DataImportValidationError(f"第 {row_number} 行包含无表头的额外列")
            result.append((row_number, dict(zip(headers, values, strict=True))))
        return headers, result

    def _logical_records(
        self,
        kind: DataImportType,
        rows: list[tuple[int, dict[str, object]]],
        mapping: dict[str, str],
    ) -> list[tuple[int, dict[str, object]]]:
        if kind is not DataImportType.ORDER:
            return rows
        order_header = mapping["external_order_id"]
        grouped: dict[str, list[tuple[int, dict[str, object]]]] = defaultdict(list)
        for row_number, row in rows:
            identity = str(row.get(order_header, "")) or f"__invalid_row_{row_number}"
            grouped[identity].append((row_number, row))
        return [(items[0][0], {"rows": [row for _, row in items]}) for items in grouped.values()]

    @staticmethod
    def _apply_mapping(raw: object, mapping: dict[str, str]) -> dict[str, str]:
        if not isinstance(raw, dict):
            raise DataImportValidationError("导入行结构无效")
        return {field: str(raw.get(header, "")) for field, header in mapping.items()}

    @staticmethod
    def _normalize_mapping(kind: DataImportType, mapping: dict[str, str]) -> dict[str, str]:
        required, optional = FIELD_SETS[kind]
        allowed = required | optional
        normalized: dict[str, str] = {}
        for field, header in mapping.items():
            if field not in allowed:
                raise DataImportValidationError(f"不支持的映射字段: {field}")
            if not isinstance(header, str) or not header.strip():
                raise DataImportValidationError("映射目标表头不能为空")
            normalized[field] = header.strip()
        if len(set(normalized.values())) != len(normalized):
            raise DataImportValidationError("多个逻辑字段不得映射到同一列")
        return dict(sorted(normalized.items()))

    @staticmethod
    def _resolve_mapping(
        kind: DataImportType, mapping: dict[str, str], headers: list[str]
    ) -> dict[str, str]:
        if any(SENSITIVE_FIELD_PATTERN.search(header) for header in headers):
            raise DataImportValidationError("导入表头包含敏感字段")
        required, optional = FIELD_SETS[kind]
        resolved = dict(mapping)
        for field in required | optional:
            if field not in resolved and field in headers:
                resolved[field] = field
        missing = sorted(required - resolved.keys())
        if missing:
            raise DataImportValidationError(f"缺少必需列映射: {', '.join(missing)}")
        absent = sorted({header for header in resolved.values() if header not in headers})
        if absent:
            raise DataImportValidationError(f"映射列不存在: {', '.join(absent)}")
        if set(headers) != set(resolved.values()):
            raise DataImportValidationError("导入文件包含未映射列")
        return dict(sorted(resolved.items()))

    @staticmethod
    def _event_type(
        kind: DataImportType, raw_values: dict[str, object], mapping: dict[str, str]
    ) -> str:
        if kind is DataImportType.ORDER:
            return "ORDER.SNAPSHOT"
        if kind is DataImportType.INVENTORY:
            inventory_header = mapping.get("inventory_kind", "inventory_kind")
            value = str(raw_values.get(inventory_header, "")).strip().upper()
            if value == "CHANNEL":
                return "INVENTORY.CHANNEL_SNAPSHOT"
            return "INVENTORY.WAREHOUSE_SNAPSHOT"
        return f"IMPORT.{kind.value}"

    def _source_time(
        self, kind: DataImportType, raw_values: dict[str, object], mapping: dict[str, str]
    ) -> datetime | None:
        raw: object = raw_values
        if kind is DataImportType.ORDER:
            rows = raw_values.get("rows")
            if not isinstance(rows, list) or not rows:
                return None
            raw = rows[0]
            field = "source_updated_at" if "source_updated_at" in mapping else "ordered_at"
        elif kind is DataImportType.INVENTORY:
            field = "source_updated_at"
        elif kind is DataImportType.COST:
            field = "effective_from"
        else:
            return None
        if not isinstance(raw, dict):
            return None
        value = str(raw.get(mapping.get(field, field), ""))
        try:
            return self._aware_datetime(value, field)
        except DataImportValidationError:
            return None

    @staticmethod
    def _aware_datetime(value: str, label: str) -> datetime:
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError as exc:
            raise DataImportValidationError(f"{label} 不是有效时间") from exc
        if parsed.tzinfo is None or parsed.utcoffset() is None:
            raise DataImportValidationError(f"{label} 必须包含时区")
        return parsed.astimezone(UTC)

    @staticmethod
    def _validation_errors(exc: Exception) -> list[dict[str, object]]:
        if isinstance(exc, ValidationError):
            return [
                {
                    "field": ".".join(str(item) for item in error.get("loc", ())),
                    "code": str(error.get("type", "VALIDATION_ERROR")),
                    "message": str(error.get("msg", "字段无效"))[:500],
                }
                for error in exc.errors()
            ]
        return [{"code": "VALIDATION_ERROR", "message": str(exc)[:500]}]

    @staticmethod
    def _import_type(value: str) -> DataImportType:
        try:
            return DataImportType(value.strip().upper())
        except ValueError as exc:
            raise DataImportValidationError(
                "import_type 必须为 CATALOG、ORDER、INVENTORY 或 COST"
            ) from exc

    @staticmethod
    def _file_identity(filename: str) -> tuple[str, str]:
        safe_name = Path(filename.replace("\\", "/")).name.strip()
        if not safe_name or len(safe_name) > 255:
            raise DataImportValidationError("文件名无效")
        extension = Path(safe_name).suffix.lower()
        if extension not in {".csv", ".xlsx"}:
            raise DataImportValidationError("仅支持 .csv 和 .xlsx 文件")
        return safe_name, extension[1:]

    @staticmethod
    def _optional_text(value: object) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    def _sku_by_code(self, sku_code: str) -> MasterSKU:
        sku = self.session.scalar(
            select(MasterSKU).where(
                MasterSKU.organization_id == self.principal.organization_id,
                MasterSKU.sku_code == sku_code.strip().upper(),
                MasterSKU.active.is_(True),
            )
        )
        if sku is None:
            raise DataImportValidationError("主 SKU 编码不存在或不可用")
        return sku

    def _job(self, job_id: int, *, for_update: bool = False) -> DataImportJob:
        statement = select(DataImportJob).where(
            DataImportJob.id == job_id,
            DataImportJob.organization_id == self.principal.organization_id,
        )
        if for_update:
            statement = statement.with_for_update().execution_options(populate_existing=True)
        job = self.session.scalar(statement)
        if job is None:
            raise DataImportNotFoundError("导入任务不存在")
        return job

    def _record(self, record_id: int, *, for_update: bool = False) -> DataImportRecord:
        statement = select(DataImportRecord).where(
            DataImportRecord.id == record_id,
            DataImportRecord.organization_id == self.principal.organization_id,
        )
        if for_update:
            statement = statement.with_for_update().execution_options(populate_existing=True)
        record = self.session.scalar(statement)
        if record is None:
            raise DataImportNotFoundError("导入记录不存在")
        return record

    def _records(self, job_id: int) -> list[DataImportRecord]:
        return list(
            self.session.scalars(
                select(DataImportRecord)
                .where(
                    DataImportRecord.organization_id == self.principal.organization_id,
                    DataImportRecord.import_job_id == job_id,
                )
                .order_by(DataImportRecord.id)
            )
        )

    def _audit(self, tool_name: str, details: dict[str, object]) -> None:
        self.session.add(
            OperationLog(
                request_id=str(uuid4()),
                session_id=None,
                tool_name=tool_name,
                tool_input={
                    "actor_user_id": self.principal.user_id,
                    "organization_id": self.principal.organization_id,
                    **details,
                },
                tool_output={"status": "SUCCESS"},
                duration_ms=0,
                status="SUCCESS",
            )
        )
